from pathlib import Path

from docx import Document as WordDocument
from pptx import Presentation


def extract_docx(path: str) -> dict:
    """
    Extract text from a Microsoft Word DOCX file.
    """

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"DOCX file not found: {file_path}"
        )

    document = WordDocument(file_path)

    parts = []

    # Paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            parts.append(text)

    # Tables
    for table_number, table in enumerate(
        document.tables,
        start=1
    ):
        parts.append(
            f"[Table {table_number}]"
        )

        for row in table.rows:
            values = []

            for cell in row.cells:
                values.append(
                    cell.text.strip()
                )

            parts.append(
                " | ".join(values)
            )

    text = "\n".join(parts).strip()

    return {
        "text": text,
        "ocr_used": False,
        "page_count": 0,
        "file_type": "docx",
    }


def extract_pptx(path: str) -> dict:
    """
    Extract text from a Microsoft PowerPoint PPTX file.
    """

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"PPTX file not found: {file_path}"
        )

    presentation = Presentation(file_path)

    parts = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_parts = [
            f"[Slide {slide_number}]"
        ]

        for shape in slide.shapes:

            # Normal text boxes
            if hasattr(shape, "text"):
                text = shape.text.strip()

                if text:
                    slide_parts.append(text)

            # Tables inside slides
            if getattr(shape, "has_table", False):

                table = shape.table

                for row in table.rows:
                    values = []

                    for cell in row.cells:
                        values.append(
                            cell.text.strip()
                        )

                    slide_parts.append(
                        " | ".join(values)
                    )

        parts.append(
            "\n".join(slide_parts)
        )

    text = "\n\n".join(parts).strip()

    return {
        "text": text,
        "ocr_used": False,
        "page_count": len(presentation.slides),
        "file_type": "pptx",
    }
