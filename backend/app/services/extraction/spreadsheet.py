from pathlib import Path
import csv

import openpyxl

try:  # .xls support is optional
    import xlrd
except ImportError:  # pragma: no cover
    xlrd = None


def extract_xlsx(path: str) -> dict:
    """
    Extract values from an XLSX workbook.
    """

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"XLSX file not found: {file_path}"
        )

    workbook = openpyxl.load_workbook(
        file_path,
        read_only=True,
        data_only=True
    )

    parts = []

    for worksheet in workbook.worksheets:

        parts.append(
            f"[Sheet: {worksheet.title}]"
        )

        for row in worksheet.iter_rows(
            values_only=True
        ):

            values = []

            for value in row:

                if value is None:
                    values.append("")
                else:
                    values.append(
                        str(value)
                    )

            # Ignore completely empty rows
            if any(value.strip() for value in values):
                parts.append(
                    " | ".join(values)
                )

    workbook.close()

    text = "\n".join(parts).strip()

    return {
        "text": text,
        "ocr_used": False,
        "page_count": len(workbook.sheetnames),
        "file_type": "xlsx",
    }


def extract_xls(path: str) -> dict:
    """
    Extract values from legacy XLS files.
    """

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"XLS file not found: {file_path}"
        )

    if xlrd is None:
        raise RuntimeError("Reading .xls files needs xlrd. Run: pip install xlrd")

    workbook = xlrd.open_workbook(
        file_path
    )

    parts = []

    for sheet in workbook.sheets():

        parts.append(
            f"[Sheet: {sheet.name}]"
        )

        for row_index in range(
            sheet.nrows
        ):

            values = []

            for column_index in range(
                sheet.ncols
            ):

                value = sheet.cell_value(
                    row_index,
                    column_index
                )

                values.append(
                    str(value)
                )

            if any(
                value.strip()
                for value in values
            ):
                parts.append(
                    " | ".join(values)
                )

    text = "\n".join(parts).strip()

    return {
        "text": text,
        "ocr_used": False,
        "page_count": workbook.nsheets,
        "file_type": "xls",
    }


def extract_csv(path: str) -> dict:
    """
    Extract CSV rows into searchable text.
    """

    file_path = Path(path)

    if not file_path.exists():
        raise FileNotFoundError(
            f"CSV file not found: {file_path}"
        )

    encodings = [
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1",
    ]

    rows = None

    for encoding in encodings:

        try:

            with open(
                file_path,
                "r",
                encoding=encoding,
                newline=""
            ) as file:

                rows = list(
                    csv.reader(file)
                )

            break

        except UnicodeDecodeError:
            continue

    if rows is None:
        raise RuntimeError(
            "Unable to decode CSV file."
        )

    parts = [
        "[CSV DATA]"
    ]

    for row in rows:

        values = [
            str(value).strip()
            for value in row
        ]

        if any(values):
            parts.append(
                " | ".join(values)
            )

    text = "\n".join(parts).strip()

    return {
        "text": text,
        "ocr_used": False,
        "page_count": 1,
        "file_type": "csv",
    }
